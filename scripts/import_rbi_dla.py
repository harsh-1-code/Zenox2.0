#!/usr/bin/env python3
"""Load the RBI Digital Lending Apps directory into the app's reference list.

    python scripts/import_rbi_dla.py ~/Downloads/dla_directory.xlsx

Where to get the file:
    rbi.org.in -> Citizen's Corner -> "DLA's deployed by Regulated Entities"
    -> export the report to CSV or Excel from the viewer's toolbar.

The directory is served through a JavaScript report viewer, so it cannot be fetched
programmatically - the manual export is a property of the source, not a shortcut here.
Column names vary between exports, so headers are matched loosely.
"""

import json
import sys
from datetime import date
from pathlib import Path

DATA = Path(__file__).resolve().parent.parent / "backend/data/lending_apps.json"

APP_COLS = ("dla", "app name", "name of dla", "digital lending app", "app")
ENT_COLS = ("regulated entity", "name of re", "re name", "entity", "lender")
TYPE_COLS = ("entity type", "re type", "category", "type")


def pick(headers: list, wanted: tuple) -> int:
    for i, h in enumerate(headers):
        if any(w == str(h).strip().lower() for w in wanted):
            return i
    for i, h in enumerate(headers):  # substring fallback
        if any(w in str(h).strip().lower() for w in wanted):
            return i
    return -1


def read_rows(path: Path) -> list:
    if path.suffix.lower() in {".xlsx", ".xls"}:
        try:
            from openpyxl import load_workbook
        except ImportError:
            sys.exit("pip install openpyxl, or export the directory as CSV instead")
        ws = load_workbook(path, read_only=True, data_only=True).active
        return [list(row) for row in ws.iter_rows(values_only=True)]
    import csv

    with path.open(newline="", encoding="utf-8-sig") as f:
        return list(csv.reader(f))


def main() -> None:
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    rows = [r for r in read_rows(Path(sys.argv[1])) if r and any(c for c in r)]
    if not rows:
        sys.exit("file is empty")

    # The export usually carries title rows above the real header.
    head_i = next((i for i, r in enumerate(rows[:12]) if pick(list(r), APP_COLS) >= 0), -1)
    if head_i < 0:
        sys.exit(f"no app-name column found. First row seen: {rows[0]}")

    headers = list(rows[head_i])
    ai, ei, ti = pick(headers, APP_COLS), pick(headers, ENT_COLS), pick(headers, TYPE_COLS)

    def cell(r: list, i: int) -> str:
        return str(r[i]).strip() if 0 <= i < len(r) and r[i] is not None else ""

    apps, seen = [], set()
    for r in rows[head_i + 1 :]:
        app = cell(r, ai)
        if not app or app.lower() in seen:
            continue
        seen.add(app.lower())
        apps.append({
            "app": app,
            "entity": cell(r, ei) or "Not stated",
            "entity_type": cell(r, ti) or "RE",
        })

    cur = json.loads(DATA.read_text())
    cur["apps"] = apps
    cur["last_updated"] = date.today().isoformat()
    cur["_status"] = "LOADED"
    cur.pop("_why", None)
    DATA.write_text(json.dumps(cur, ensure_ascii=False, indent=2) + "\n")
    print(f"imported {len(apps)} apps -> {DATA}")
    print(f"last_updated = {cur['last_updated']} (shown to the user, as the PS requires)")


if __name__ == "__main__":
    main()
